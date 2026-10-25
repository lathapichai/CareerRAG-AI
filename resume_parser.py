from pathlib import Path
from io import BytesIO

import fitz
from docx import Document


def extract_text_from_pdf(file):
    try:
        if isinstance(file, (str, Path)):
            pdf = fitz.open(str(file))
        else:
            content = file.read()
            pdf = fitz.open(stream=content, filetype="pdf")

        try:
            text = "\n".join(page.get_text() for page in pdf)
        finally:
            pdf.close()

        return text.strip()

    except Exception as error:
        raise ValueError(f"Could not read PDF resume: {error}")


def extract_text_from_docx(file):
    try:
        if isinstance(file, (str, Path)):
            document = Document(str(file))
        else:
            content = file.read()
            document = Document(BytesIO(content))

        return "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ).strip()

    except Exception as error:
        raise ValueError(f"Could not read DOCX resume: {error}")
